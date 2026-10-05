from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SOURCE = Path(r"E:\dynamicmart\docs\report-edit\Bao-cao-nguon.docx")
OUTPUT = Path(r"E:\dynamicmart\docs\report-edit\Bao-cao-DynamicMart-dong-bo.docx")


def normalized(value: str) -> str:
    return " ".join((value or "").split())


def find_paragraph(doc: Document, text: str):
    target = normalized(text)
    for paragraph in doc.paragraphs:
        if normalized(paragraph.text) == target:
            return paragraph
    raise ValueError(f"Không tìm thấy đoạn: {text}")


def replace_text(paragraph, text: str):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def delete_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def clear_between(start_paragraph, end_paragraph):
    current = start_paragraph._p.getnext()
    while current is not None and current is not end_paragraph._p:
        following = current.getnext()
        delete_element(current)
        current = following


def add_paragraph_after(doc, anchor, text: str, style: str = "Normal", bold: bool = False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold = bold
    anchor.addnext(paragraph._p)
    return paragraph._p


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def format_table(table, widths=None):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    fitted_widths = None
    if widths:
        usable_width = 6.20
        scale = min(1.0, usable_width / sum(widths))
        fitted_widths = [value * scale for value in widths]
        grid = table._tbl.tblGrid
        grid_cols = list(grid.gridCol_lst)
        for idx, value in enumerate(fitted_widths):
            if idx < len(grid_cols):
                grid_cols[idx].set(qn("w:w"), str(int(value * 1440)))
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if fitted_widths and col_idx < len(fitted_widths):
                cell.width = Inches(fitted_widths[col_idx])
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(int(fitted_widths[col_idx] * 1440)))
                tc_w.set(qn("w:type"), "dxa")
            if row_idx == 0:
                shade_cell(cell, "D9E2F3")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if row_idx == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(10)
                    if row_idx == 0:
                        run.bold = True


def add_table_after(doc, anchor, rows, widths=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            table.cell(r_idx, c_idx).text = str(value)
    format_table(table, widths)
    anchor.addnext(table._tbl)
    return table._tbl


def replace_section(doc, start_text: str, end_text: str, builder):
    start = find_paragraph(doc, start_text)
    end = find_paragraph(doc, end_text)
    clear_between(start, end)
    builder(start._p)


doc = Document(SOURCE)

# Mở đầu: dùng giọng văn trung tính, không thuật lại vai trò của người viết.
replacements = {
    "Trước hết, nhóm tiến hành khảo sát yêu cầu của một hệ thống thương mại điện tử và xác định các chức năng chính cần xây dựng. Từ đó, hệ thống được phân chia thành các nhóm nghiệp vụ phù hợp để thuận lợi cho việc thiết kế và phân công công việc.":
        "Quá trình thực hiện bắt đầu bằng việc khảo sát yêu cầu của hệ thống thương mại điện tử và xác định phạm vi nghiệp vụ. Ba use case trọng tâm gồm Danh mục sản phẩm, Đơn hàng và Thanh toán; các chức năng tài khoản, giỏ hàng, voucher và vận chuyển đóng vai trò hỗ trợ cho ba luồng này.",
    "Tiếp theo, nhóm thiết kế giao diện trao đổi dữ liệu, cấu trúc hệ thống và các quy tắc xử lý cơ bản trước khi tiến hành lập trình. Các thành viên phát triển từng phần theo nhiệm vụ được phân công và thống nhất cách tích hợp giữa các thành phần.":
        "Kiến trúc, hợp đồng API, mô hình dữ liệu và quy tắc nghiệp vụ được xác định trước khi hiện thực. Các thành phần được phát triển theo ranh giới dịch vụ và tích hợp thông qua API Gateway hoặc sự kiện.",
    "Cuối cùng, nhóm tiến hành kiểm thử, sửa lỗi và đối chiếu kết quả với yêu cầu ban đầu để đánh giá mức độ hoàn thiện của hệ thống.":
        "Cuối cùng, các API, giao diện và luồng nghiệp vụ được kiểm thử, sửa lỗi và đối chiếu với yêu cầu ban đầu để đánh giá mức độ hoàn thiện.",
    "2.1.4 Yêu cầu phi chức năng": "2.1.4 Quy tắc nghiệp vụ",
    "D) Chống xử lý trùng lặp": "d) Chống xử lý trùng lặp",
}
for old, new in replacements.items():
    try:
        replace_text(find_paragraph(doc, old), new)
    except ValueError:
        pass


# 2.1.2 Dữ liệu cần quản lý
def build_data_scope(anchor):
    anchor = add_paragraph_after(doc, anchor, "Dữ liệu được tổ chức theo ba phạm vi nghiệp vụ chính và một số dữ liệu hỗ trợ:")
    anchor = add_paragraph_after(doc, anchor, "- Danh mục sản phẩm: danh mục, thuộc tính, sản phẩm, biến thể, hình ảnh, giá bán, tồn kho và nguồn tri thức phục vụ trợ lý mua sắm.")
    anchor = add_paragraph_after(doc, anchor, "- Đơn hàng: giỏ hàng được chọn, phiên đặt hàng, địa chỉ giao hàng, voucher, phí vận chuyển, các dòng hàng và lịch sử trạng thái.")
    anchor = add_paragraph_after(doc, anchor, "- Thanh toán: nghĩa vụ thanh toán, các lần thực hiện, kết quả phản hồi, phương thức COD hoặc VNPay và trạng thái giao dịch.")
    anchor = add_paragraph_after(doc, anchor, "- Dữ liệu hỗ trợ: tài khoản, phiên đăng nhập, địa chỉ cá nhân và các bản ghi sự kiện cần thiết cho xác thực, truy vết và tích hợp.")
    add_paragraph_after(doc, anchor, "Thông tin đã ghi nhận trong đơn hàng được lưu dưới dạng bản chụp. Việc thay đổi sản phẩm, giá bán hoặc địa chỉ sau đó không làm thay đổi nội dung giao dịch đã phát sinh.")


replace_section(doc, "2.1.2 Dữ liệu cần quản lý", "2.1.3. Yêu cầu chức năng trọng tâm", build_data_scope)


# 2.1.3 Ba use case trọng tâm
def build_function_scope(anchor):
    anchor = add_paragraph_after(doc, anchor, "Phạm vi phân tích và kiểm chứng tập trung vào ba use case chính:")
    anchor = add_paragraph_after(doc, anchor, "- Danh mục sản phẩm: quản lý danh mục, sản phẩm, biến thể, hình ảnh và tồn kho; hỗ trợ duyệt, tìm kiếm, lọc, xem chi tiết và hỏi đáp qua trợ lý mua sắm.")
    anchor = add_paragraph_after(doc, anchor, "- Đơn hàng: chuẩn bị dữ liệu đặt hàng, kiểm tra địa chỉ, voucher, vận chuyển và tồn kho; tạo đơn, lưu bản chụp và theo dõi trạng thái.")
    anchor = add_paragraph_after(doc, anchor, "- Thanh toán: lựa chọn COD hoặc VNPay, khởi tạo giao dịch, xác minh kết quả và cung cấp trạng thái thanh toán cho đơn hàng.")
    add_paragraph_after(doc, anchor, "Tài khoản, giỏ hàng, voucher và vận chuyển không được tách thành use case trọng tâm; các chức năng này cung cấp dữ liệu và điều kiện cần thiết cho quá trình đặt hàng và thanh toán.")


replace_section(doc, "2.1.3. Yêu cầu chức năng trọng tâm", "2.1.4 Quy tắc nghiệp vụ", build_function_scope)

# Bổ sung quy tắc cho trợ lý vào cuối mục 2.1.4.
rule_end = find_paragraph(doc, "Những quy tắc trên thể hiện phần logic nghiệp vụ cốt lõi của DynamicMart, giúp hệ thống không chỉ dừng lại ở các thao tác thêm, sửa, xóa dữ liệu.")
replace_text(rule_end, "Các quy tắc trên tạo thành logic nghiệp vụ cốt lõi. Trợ lý mua sắm chỉ trả lời từ nguồn tri thức đã kiểm duyệt, không tự tính giá, tồn kho, tổng tiền, trạng thái đơn hoặc trạng thái thanh toán; các giá trị quyết định luôn được lấy từ API nghiệp vụ.")


# 2.2 Mô hình hóa ba use case
def build_use_cases(anchor):
    anchor = add_paragraph_after(doc, anchor, "Phần mô hình hóa chỉ sử dụng ba use case trọng tâm. Những thao tác hỗ trợ được đặt trong luồng của use case tương ứng để tránh chia nhỏ nghiệp vụ thành danh sách CRUD.")
    anchor = add_paragraph_after(doc, anchor, "2.2.1 Sơ đồ Use Case", "Heading 3")
    anchor = add_paragraph_after(doc, anchor, "Sơ đồ thể hiện khách hàng và quản trị viên tương tác với ba use case Danh mục sản phẩm, Đơn hàng và Thanh toán. Cổng thanh toán và đơn vị vận chuyển chỉ tham gia tại các bước tích hợp tương ứng.")
    anchor = add_paragraph_after(doc, anchor, "[CHÈN HÌNH 2.1: Sơ đồ Use Case tổng quát gồm Danh mục sản phẩm, Đơn hàng và Thanh toán]")
    anchor = add_paragraph_after(doc, anchor, "2.2.2 Đặc tả Use Case", "Heading 3")
    anchor = add_paragraph_after(doc, anchor, "Ba use case được đặc tả ở mức đủ để xác định tác nhân, điều kiện, luồng xử lý và kết quả cần kiểm chứng.")
    rows = [
        ["Mã và Use Case", "Tác nhân và điều kiện", "Luồng chính", "Ngoại lệ và kết quả"],
        ["UC01\nDanh mục sản phẩm", "Khách truy cập, khách hàng, quản trị viên. Dữ liệu danh mục đã được công bố.", "Duyệt hoặc quản lý danh mục; tìm kiếm, lọc, chọn biến thể, xem giá và tồn kho. Trợ lý hỗ trợ giải đáp theo nguồn đã kiểm duyệt.", "Điều kiện lọc không hợp lệ hoặc không có dữ liệu được phản hồi rõ. Kết quả hiển thị đúng dữ liệu hiện hành; trợ lý không suy đoán khi thiếu nguồn."],
        ["UC02\nĐơn hàng", "Khách hàng đã đăng nhập; mặt hàng, địa chỉ và phiên đặt hàng còn hợp lệ.", "Tạo phiên, chọn địa chỉ và voucher, lấy bản xem trước, kiểm tra lại dữ liệu, giữ tài nguyên và tạo đơn bằng khóa chống gửi trùng.", "Hết hàng, voucher hoặc báo giá hết hạn làm dừng luồng và hoàn phần đã giữ. Một yêu cầu hợp lệ chỉ tạo một đơn hàng."],
        ["UC03\nThanh toán", "Đơn hàng đã tạo và có nghĩa vụ thanh toán hợp lệ.", "Chọn COD hoặc VNPay; khởi tạo lần thanh toán; xác minh phản hồi và truy vấn trạng thái theo đơn.", "Sai chữ ký, sai số tiền hoặc phản hồi trùng không làm thay đổi giao dịch sai lệch. Kết quả hợp lệ được ghi nhận đúng một lần."],
    ]
    anchor = add_table_after(doc, anchor, rows, [1.1, 1.7, 2.25, 2.25])
    anchor = add_paragraph_after(doc, anchor, "2.2.3 Biểu đồ hoạt động và biểu đồ tuần tự", "Heading 3")
    anchor = add_paragraph_after(doc, anchor, "Các biểu đồ chi tiết bám theo đúng ba use case trên và sử dụng cùng tên gọi trong toàn bộ báo cáo.")
    anchor = add_paragraph_after(doc, anchor, "[CHÈN HÌNH 2.2: Biểu đồ hoạt động và tuần tự của Use Case Danh mục sản phẩm]")
    anchor = add_paragraph_after(doc, anchor, "[CHÈN HÌNH 2.3: Biểu đồ hoạt động và tuần tự của Use Case Đơn hàng]")
    add_paragraph_after(doc, anchor, "[CHÈN HÌNH 2.4: Biểu đồ hoạt động và tuần tự của Use Case Thanh toán]")


replace_section(doc, "2.2. Mô hình hóa chức năng và nghiệp vụ", "2.3 Thiết kế kiến trúc hệ thống", build_use_cases)


# 2.3: cập nhật thành phần trợ lý AI.
arch_replacements = {
    "Kiến trúc của DynamicMart gồm năm nhóm thành phần chính: ứng dụng khách, API Gateway, các dịch vụ nghiệp vụ, hệ thống lưu trữ dữ liệu và các dịch vụ bên ngoài.":
        "Kiến trúc DynamicMart gồm ứng dụng khách, API Gateway, sáu dịch vụ nghiệp vụ, dịch vụ trợ lý AI, các kho dữ liệu và những hệ thống tích hợp bên ngoài.",
    "Phần Backend gồm sáu dịch vụ nghiệp vụ:": "Phần nghiệp vụ cốt lõi gồm sáu dịch vụ, bên cạnh một dịch vụ trợ lý AI hỗ trợ tra cứu:",
    "RabbitMQ được sử dụng cho các thông tin không yêu cầu phản hồi ngay, chẳng hạn thông báo về nghĩa vụ thanh toán hoặc kết quả của một giao dịch. Ngoài ra, DynamicMart kết nối với cổng thanh toán và dịch vụ GHN để hỗ trợ thanh toán trực tuyến, tra cứu địa giới và tính phí giao hàng.":
        "RabbitMQ truyền các sự kiện không yêu cầu phản hồi tức thời. DynamicMart còn kết nối với GHN, cổng thanh toán và dịch vụ trợ lý AI. Trợ lý sử dụng FastAPI, Ollama và Qdrant để truy xuất nguồn tri thức; mọi dữ liệu quyết định về sản phẩm, đơn hàng và thanh toán vẫn do các dịch vụ nghiệp vụ cung cấp.",
}
for old, new in arch_replacements.items():
    try:
        replace_text(find_paragraph(doc, old), new)
    except ValueError:
        pass

engagement_line = find_paragraph(doc, "Quản lý đánh giá, danh sách yêu thích và thông báo.")
add_paragraph_after(doc, engagement_line._p, "Hỗ trợ hỏi đáp bằng RAG từ nguồn tri thức đã kiểm duyệt.", "List Paragraph")

# Chương 1 chỉ giới thiệu vai trò công nghệ, không lặp lại chi tiết cài đặt.
for table in doc.tables:
    if table.rows and normalized(table.cell(0, 0).text) == "Thành phần" and normalized(table.cell(0, 1).text) == "Công nghệ" and any("Spring Boot" in row.cells[1].text for row in table.rows[1:]):
        existing = {normalized(row.cells[0].text) for row in table.rows[1:]}
        if "Trợ lý AI" not in existing:
            row = table.add_row()
            row.cells[0].text = "Trợ lý AI"
            row.cells[1].text = "Python, FastAPI, Ollama"
            row.cells[2].text = "Cung cấp API hỏi đáp và sinh câu trả lời từ ngữ cảnh truy xuất"
        if "Cơ sở dữ liệu vector" not in existing:
            row = table.add_row()
            row.cells[0].text = "Cơ sở dữ liệu vector"
            row.cells[1].text = "Qdrant"
            row.cells[2].text = "Lưu vector và truy xuất nguồn tri thức theo tenant"
        format_table(table, [1.55, 2.15, 3.5])
        break

# Thêm dòng vào bảng thành phần kiến trúc.
for table in doc.tables:
    if table.rows and normalized(table.cell(0, 0).text) == "Thành phần" and any("Identity Service" in row.cells[0].text for row in table.rows[1:]):
        row = table.add_row()
        row.cells[0].text = "AI Assistant Service"
        row.cells[1].text = "Tiếp nhận câu hỏi, truy xuất nguồn tri thức theo tenant và trả lời kèm trích dẫn; không thay thế dữ liệu nghiệp vụ của Catalog, Order hoặc Payment."
        format_table(table, [1.55, 5.8])
        break


# 2.4.1 và 2.4.2 chỉ giữ dữ liệu cốt lõi của ba use case.
def build_data_design(anchor):
    anchor = add_paragraph_after(doc, anchor, "Ba nhóm dữ liệu dưới đây tương ứng trực tiếp với ba use case trọng tâm. Tài khoản, giỏ hàng và vận chuyển được tham chiếu khi cần nhưng không mở rộng thành nhóm thiết kế riêng trong phần này.")
    rows = [
        ["Nhóm dữ liệu", "Thực thể cốt lõi", "Nội dung quản lý"],
        ["Danh mục sản phẩm", "Category, Attribute, Product, ProductVariant, ProductImage, InventoryItem", "Cấu trúc danh mục, thông tin bán, biến thể, hình ảnh và số lượng tồn kho"],
        ["Đơn hàng", "CheckoutSession, Order, OrderItem, OrderAddress, OrderStatusHistory", "Phiên đặt hàng, bản chụp sản phẩm và địa chỉ, tổng tiền và vòng đời đơn"],
        ["Thanh toán", "Payment, PaymentAttempt, PaymentCallbackAudit", "Nghĩa vụ thanh toán, các lần thực hiện và kết quả phản hồi"],
    ]
    add_table_after(doc, anchor, rows, [1.35, 3.15, 2.7])


replace_section(doc, "2.4.1. Các nhóm thực thể chính", "2.4.2. Các thực thể và quan hệ cốt lõi", build_data_design)


def build_core_entities(anchor):
    rows = [
        ["Thực thể", "Khóa chính", "Khóa ngoại hoặc tham chiếu", "Quan hệ chính"],
        ["categories", "id", "parent_id → categories.id", "Một danh mục có thể có nhiều danh mục con và nhiều sản phẩm"],
        ["products", "id", "category_id → categories.id", "Một sản phẩm thuộc một danh mục và có nhiều biến thể"],
        ["product_variants", "id", "product_id → products.id", "Mỗi biến thể thuộc một sản phẩm và có tối đa một bản ghi tồn kho"],
        ["inventory_items", "variant_id", "variant_id → product_variants.id", "Lưu số lượng thực tế và số lượng đang giữ của biến thể"],
        ["checkout_sessions", "id", "customer_id, cart_id, address_id là REF", "Một phiên đặt hàng tạo tối đa một đơn hàng"],
        ["orders", "id", "checkout_session_id UNIQUE; customer_id REF", "Một đơn có nhiều dòng hàng, một địa chỉ và nhiều trạng thái"],
        ["order_items", "id", "order_id → orders.id; product_id, variant_id là REF", "Lưu bản chụp sản phẩm tại thời điểm đặt hàng"],
        ["order_addresses", "id", "order_id → orders.id", "Mỗi đơn lưu cố định một địa chỉ giao hàng"],
        ["payments", "id", "order_id REF; customer_id REF", "Mỗi đơn có tối đa một nghĩa vụ thanh toán"],
        ["payment_attempts", "id", "payment_id → payments.id", "Một khoản thanh toán có thể có nhiều lần thực hiện"],
        ["payment_callback_audits", "id", "payment_attempt_id → payment_attempts.id", "Lưu kết quả xác minh phản hồi thanh toán"],
    ]
    add_table_after(doc, anchor, rows, [1.35, 1.05, 2.55, 2.35])


replace_section(doc, "2.4.2. Các thực thể và quan hệ cốt lõi", "2.4.3. Sơ đồ ERD", build_core_entities)


# 2.5: danh sách API gọn theo ba use case và trợ lý hỗ trợ Catalog.
def build_api_design(anchor):
    anchor = add_paragraph_after(doc, anchor, "Thiết kế API bám theo ba tài nguyên nghiệp vụ chính. Các endpoint hỗ trợ được đặt trong luồng tương ứng để tránh lặp lại danh sách chức năng ở nhiều chương.")
    anchor = add_paragraph_after(doc, anchor, "2.5.1 Quy ước API", "Heading 3")
    rules = [
        ["Nội dung", "Quy ước"],
        ["Địa chỉ chung", "http://localhost:8080/api/v1 qua API Gateway"],
        ["Giao thức và dữ liệu", "HTTP hoặc HTTPS; request và response sử dụng JSON"],
        ["Xác thực", "JWT trong Authorization; refresh token lưu bằng cookie bảo mật"],
        ["Tên tài nguyên", "Danh từ viết thường, số nhiều; tham số đường dẫn xác định tài nguyên"],
        ["Tìm kiếm", "Query parameter dùng cho từ khóa, lọc, sắp xếp và phân trang"],
        ["Chống gửi trùng", "Idempotency-Key bắt buộc khi tạo đơn hàng"],
        ["API nội bộ", "Dùng khóa nội bộ và không cho ứng dụng khách gọi trực tiếp"],
    ]
    anchor = add_table_after(doc, anchor, rules, [1.7, 5.5])
    anchor = add_paragraph_after(doc, anchor, "2.5.2 Các API chính", "Heading 3")
    apis = [
        ["Use Case", "Phương thức và endpoint", "Dữ liệu chính", "Phản hồi thành công", "Lỗi chính"],
        ["Danh mục sản phẩm", "GET /catalog/categories", "Điều kiện trạng thái", "Cây danh mục", "400, 503"],
        ["Danh mục sản phẩm", "GET /catalog/products", "Từ khóa, danh mục, giá, thuộc tính, phân trang", "Danh sách sản phẩm", "400, 503"],
        ["Danh mục sản phẩm", "GET /catalog/products/{slug}", "slug sản phẩm", "Chi tiết, biến thể, ảnh và tồn kho", "404"],
        ["Hỗ trợ danh mục", "POST /assistant/chat", "message; tenant được Gateway cố định", "Câu trả lời và nguồn trích dẫn", "400, 503"],
        ["Đơn hàng", "POST /checkout/sessions", "Giỏ hàng hoặc lựa chọn mua", "Phiên đặt hàng", "400, 409"],
        ["Đơn hàng", "POST /checkout/sessions/{id}/preview", "Địa chỉ, voucher, vận chuyển", "Bản xem trước tổng tiền", "409, 422"],
        ["Đơn hàng", "POST /checkout/orders", "Phiên đặt hàng; Idempotency-Key", "Đơn hàng đã tạo", "409, 422, 503"],
        ["Đơn hàng", "GET /orders/{orderId}", "orderId và JWT", "Chi tiết cùng lịch sử trạng thái", "401, 404"],
        ["Thanh toán", "GET /payments/orders/{orderId}", "orderId và JWT", "Thông tin và trạng thái thanh toán", "401, 404"],
        ["Thanh toán", "POST /payments/{id}/vnpay-attempt", "paymentId", "Đường dẫn thanh toán", "409, 502"],
        ["Thanh toán", "POST /payments/{id}/cod-confirmations", "paymentId và thông tin xác nhận", "Trạng thái COD", "403, 409"],
        ["Thanh toán", "GET /payments/return-status", "Mã tham chiếu giao dịch", "Kết quả đã được xác minh", "400, 404"],
    ]
    anchor = add_table_after(doc, anchor, apis, [0.95, 1.75, 1.45, 1.35, 0.70])
    anchor = add_paragraph_after(doc, anchor, "2.5.3 Phản hồi và xử lý lỗi", "Heading 3")
    anchor = add_paragraph_after(doc, anchor, "Phản hồi thành công trả về dữ liệu phù hợp với tài nguyên. Phản hồi lỗi gồm mã lỗi, thông báo và danh sách trường không hợp lệ khi cần. Ứng dụng khách sử dụng cả HTTP status code và mã lỗi nghiệp vụ để lựa chọn cách xử lý.")
    status = [
        ["Mã trạng thái", "Trường hợp sử dụng"],
        ["200 hoặc 201", "Truy vấn hoặc tạo dữ liệu thành công"],
        ["400", "Thiếu dữ liệu, sai định dạng hoặc tham số không hợp lệ"],
        ["401 hoặc 403", "Chưa xác thực hoặc không có quyền"],
        ["404", "Không tìm thấy tài nguyên thuộc phạm vi truy cập"],
        ["409 hoặc 422", "Xung đột trạng thái hoặc vi phạm quy tắc nghiệp vụ"],
        ["502 hoặc 503", "Hệ thống tích hợp hoặc dịch vụ phụ thuộc chưa sẵn sàng"],
    ]
    add_table_after(doc, anchor, status, [1.55, 5.65])


replace_section(doc, "2.5 Thiết kế API dịch vụ", "2.6 Thiết kế ứng dụng client", build_api_design)


# 2.6: màn hình và tích hợp API theo đúng ba use case.
def build_client_design(anchor):
    anchor = add_paragraph_after(doc, anchor, "Ứng dụng khách được thiết kế như thành phần tiêu thụ API qua Gateway. Giao diện chỉ kiểm tra dữ liệu đầu vào cơ bản; giá, tồn kho, tổng tiền, trạng thái đơn và trạng thái thanh toán luôn được xác minh tại máy chủ.")
    anchor = add_paragraph_after(doc, anchor, "2.6.1 Thiết kế giao diện", "Heading 3")
    screens = [
        ["Use Case", "Màn hình chính", "Nội dung"],
        ["Danh mục sản phẩm", "Danh sách sản phẩm", "Tìm kiếm, lọc, sắp xếp và phân trang"],
        ["Danh mục sản phẩm", "Chi tiết sản phẩm", "Hình ảnh, thuộc tính, biến thể, giá và khả năng mua"],
        ["Danh mục sản phẩm", "Quản trị danh mục và tồn kho", "Danh mục, sản phẩm, biến thể, trạng thái và số lượng tồn"],
        ["Danh mục sản phẩm", "Trợ lý mua sắm", "Gửi câu hỏi, nhận câu trả lời và nguồn trích dẫn"],
        ["Đơn hàng", "Xác nhận đặt hàng", "Địa chỉ, voucher, vận chuyển, phương thức và bản xem trước"],
        ["Đơn hàng", "Lịch sử và chi tiết đơn", "Mã đơn, dòng hàng, tổng tiền và lịch sử trạng thái"],
        ["Đơn hàng", "Quản trị đơn hàng", "Tra cứu và chuyển trạng thái theo quy tắc"],
        ["Thanh toán", "Lựa chọn thanh toán", "COD hoặc VNPay theo điều kiện của đơn"],
        ["Thanh toán", "Kết quả và quản trị giao dịch", "Số tiền, phương thức, lần thực hiện và trạng thái"],
    ]
    anchor = add_table_after(doc, anchor, screens, [1.35, 2.25, 3.6])
    anchor = add_paragraph_after(doc, anchor, "2.6.2 Thiết kế tích hợp API", "Heading 3")
    integrations = [
        ["Use Case", "API sử dụng", "Cách sử dụng trên giao diện"],
        ["Danh mục sản phẩm", "GET /catalog/categories; GET /catalog/products; GET /catalog/products/{slug}; POST /assistant/chat", "Tải danh mục và sản phẩm; hiển thị chi tiết; gửi câu hỏi và trình bày nguồn tham khảo"],
        ["Đơn hàng", "POST /checkout/sessions; POST /checkout/sessions/{id}/preview; POST /checkout/orders; GET /orders/{id}", "Chuẩn bị dữ liệu, tính lại tổng tiền, tạo đơn và hiển thị lịch sử"],
        ["Thanh toán", "GET /payments/orders/{orderId}; POST /payments/{id}/vnpay-attempt; POST /payments/{id}/cod-confirmations; GET /payments/return-status", "Khởi tạo phương thức phù hợp và xác minh kết quả trước khi hiển thị"],
    ]
    anchor = add_table_after(doc, anchor, integrations, [1.35, 3.15, 2.7])
    anchor = add_paragraph_after(doc, anchor, "2.6.3 Luồng dữ liệu và trạng thái", "Heading 3")
    add_paragraph_after(doc, anchor, "Thao tác trên giao diện tạo request qua lớp dùng chung, sau đó gửi đến API Gateway. Response được ánh xạ thành dữ liệu hiển thị. Mỗi màn hình phân biệt trạng thái đang tải, thành công, không có dữ liệu, lỗi nhập liệu, lỗi nghiệp vụ và lỗi kết nối. Trợ lý mua sắm hiển thị nguồn trích dẫn và thông báo chưa đủ thông tin khi không tìm thấy dữ liệu phù hợp.")


replace_section(doc, "2.6 Thiết kế ứng dụng client", "CHƯƠNG 3. CÀI ĐẶT VÀ TRIỂN KHAI HỆ THỐNG", build_client_design)


# 3.1.1: rút gọn bảng công nghệ và bổ sung công nghệ trợ lý.
def build_environment(anchor):
    rows = [
        ["Thành phần", "Công nghệ hoặc công cụ", "Phiên bản", "Mục đích sử dụng"],
        ["Hệ điều hành", "Windows", "10", "Phát triển và chạy thử"],
        ["Dịch vụ Java", "Java; Spring Boot", "17.0.12; 4.1.1", "Xây dựng Gateway và sáu dịch vụ nghiệp vụ"],
        ["Hạ tầng dịch vụ", "Spring Cloud; Maven", "2025.1.3; 3.9.16", "Định tuyến, quản lý phụ thuộc và build"],
        ["Dữ liệu nghiệp vụ", "PostgreSQL; Spring Data JPA; Flyway", "13 trở lên", "Lưu trữ, truy cập và quản lý migration"],
        ["Sự kiện", "RabbitMQ", "Theo cấu hình", "Trao đổi sự kiện bất đồng bộ"],
        ["Ứng dụng khách", "Node.js; Next.js; React; TypeScript", "22.19.0; 16.3.5; 19.2.8; 5.x", "Xây dựng giao diện và tích hợp API"],
        ["Trợ lý AI", "Python; FastAPI; Uvicorn", "3.12; 0.115.12; 0.34.3", "Cung cấp API hỏi đáp và điều phối RAG"],
        ["Mô hình ngôn ngữ", "Ollama", "qwen2.5:0.5b", "Sinh câu trả lời tiếng Việt trong môi trường cục bộ"],
        ["Mô hình embedding", "Ollama", "nomic-embed-text", "Biểu diễn nội dung tri thức thành vector"],
        ["Cơ sở dữ liệu vector", "Qdrant", "qdrant-client 1.19.1", "Lưu vector và truy xuất theo tenant"],
        ["Kiểm thử", "JUnit 5; Spring Boot Test; unittest; Postman", "Theo dự án", "Kiểm thử dịch vụ Java, trợ lý AI và hợp đồng API"],
        ["Quản lý mã nguồn", "Git", "2.43.0", "Theo dõi thay đổi"],
        ["Đóng gói", "Docker; Docker Compose", "29.7.2; 5.3.1", "Khởi tạo cơ sở dữ liệu, RabbitMQ và Qdrant"],
    ]
    add_table_after(doc, anchor, rows, [1.35, 2.15, 1.55, 2.3])


replace_section(doc, "3.1.1. Môi trường phát triển", "3.1.2. Môi trường phát triển backend", build_environment)

replace_text(find_paragraph(doc, "Backend gồm API Gateway và sáu dịch vụ nghiệp vụ:"), "Phần máy chủ gồm API Gateway, sáu dịch vụ nghiệp vụ và một dịch vụ trợ lý AI:")
for table in doc.tables:
    if len(table.columns) == 2 and table.rows and normalized(table.cell(0, 0).text) == "Thành phần" and any("Engagement Service" in row.cells[0].text for row in table.rows[1:]) and any("API Gateway" in row.cells[0].text for row in table.rows[1:]):
        if not any("AI Assistant" in row.cells[0].text for row in table.rows[1:]):
            row = table.add_row()
            row.cells[0].text = "AI Assistant Service"
            row.cells[1].text = "Hỏi đáp dựa trên RAG, trả lời kèm nguồn và tách dữ liệu theo tenant"
        format_table(table, [2.0, 5.2])


# Bổ sung hiện thực trợ lý vào API Catalog và bảng nghiệp vụ.
catalog_result = find_paragraph(doc, "Kết quả truy vấn được chuyển thành ProductSummaryResponse, chỉ chứa những dữ liệu cần thiết cho màn hình danh sách như tên sản phẩm, hình ảnh, giá bán, giá khuyến mại và trạng thái còn hàng. Việc sử dụng DTO giúp API không trả trực tiếp Entity của cơ sở dữ liệu.")
add_paragraph_after(doc, catalog_result._p, "Trợ lý mua sắm được tích hợp qua POST /api/v1/assistant/chat. Gateway chuyển yêu cầu đến dịch vụ FastAPI; nội dung được chia đoạn, tạo embedding, truy xuất trong Qdrant theo tenant và đưa vào Ollama để tạo câu trả lời kèm nguồn. Bản hiện tại chủ yếu sử dụng kho tri thức đã kiểm duyệt và chưa coi chatbot là nguồn dữ liệu giá hoặc tồn kho thời gian thực.")

for table in doc.tables:
    if table.rows and normalized(table.cell(0, 0).text) == "Nghiệp vụ" and any("Tạo đơn hàng" in row.cells[0].text for row in table.rows[1:]):
        # Xóa dòng giỏ hàng để bảng bám ba use case, sau đó thêm trợ lý hỗ trợ Catalog.
        for row in list(table.rows[1:]):
            if normalized(row.cells[0].text) == "Cập nhật giỏ hàng":
                delete_element(row._tr)
        row = table.add_row()
        row.cells[0].text = "Hỏi đáp về sản phẩm"
        row.cells[1].text = "RagService"
        row.cells[2].text = "Tenant, nguồn truy xuất, ngưỡng liên quan và tính có căn cứ của dữ liệu số"
        format_table(table, [1.65, 2.15, 3.55])
        break


# 3.3.2: hiện thực giao diện theo ba use case.
def build_client_features(anchor):
    anchor = add_paragraph_after(doc, anchor, "Các màn hình được nhóm theo ba use case trọng tâm. Những chức năng hỗ trợ chỉ xuất hiện tại bước cần sử dụng, nhờ đó cấu trúc giao diện thống nhất với phần phân tích ở Chương 2.")
    rows = [
        ["Use Case", "Màn hình hiện thực", "Chức năng chính", "API tiêu biểu"],
        ["Danh mục sản phẩm", "Danh sách, chi tiết, quản trị danh mục và tồn kho, trợ lý mua sắm", "Tìm kiếm, lọc, chọn biến thể, quản lý dữ liệu bán và hỏi đáp có nguồn", "Catalog API; POST /api/v1/assistant/chat"],
        ["Đơn hàng", "Xác nhận đặt hàng, lịch sử, chi tiết và quản trị đơn", "Tạo bản xem trước, tạo đơn bằng khóa chống gửi trùng và theo dõi trạng thái", "Checkout API; Order API"],
        ["Thanh toán", "Lựa chọn phương thức, chuyển hướng VNPay, kết quả và quản trị giao dịch", "Khởi tạo COD hoặc VNPay, kiểm tra và hiển thị trạng thái đã xác minh", "Payment API"],
    ]
    anchor = add_table_after(doc, anchor, rows, [1.25, 2.25, 2.65, 1.7])
    anchor = add_paragraph_after(doc, anchor, "Trợ lý mua sắm được gắn trong StoreLayout nên có thể mở tại các trang cửa hàng. Widget quản lý lịch sử câu hỏi trong phiên giao diện, hiển thị câu trả lời và danh sách nguồn. Khi dịch vụ không phản hồi, giao diện giữ nội dung hiện tại và cho phép gửi lại.")
    anchor = add_paragraph_after(doc, anchor, "Trong luồng Đơn hàng, mọi thay đổi về địa chỉ, voucher hoặc vận chuyển đều yêu cầu tạo lại bản xem trước. Khi tạo đơn, Idempotency-Key được giữ nguyên nếu request phải gửi lại do mất kết nối.")
    add_paragraph_after(doc, anchor, "Trong luồng Thanh toán, giao diện không suy luận kết quả từ URL trả về. Trạng thái chỉ được hiển thị sau khi Payment API xác minh và lưu kết quả.")


replace_section(doc, "3.3.2 Giao diện và chức năng chính", "3.3.3 Tích hợp Client với Backend API", build_client_features)


def build_client_integration(anchor):
    anchor = add_paragraph_after(doc, anchor, "Mỗi feature chứa kiểu dữ liệu, hàm gọi API và component riêng. apiClient gắn địa chỉ Gateway, thông tin xác thực và chuẩn hóa lỗi trước khi trả kết quả cho giao diện.")
    rows = [
        ["Use Case", "Thao tác", "API chính", "Dữ liệu hiển thị"],
        ["Danh mục sản phẩm", "Duyệt và tìm kiếm", "GET /api/v1/catalog/products", "Danh sách, giá, hình ảnh và tình trạng bán"],
        ["Danh mục sản phẩm", "Xem chi tiết", "GET /api/v1/catalog/products/{slug}", "Biến thể, thuộc tính, giá và tồn kho"],
        ["Danh mục sản phẩm", "Hỏi trợ lý", "POST /api/v1/assistant/chat", "Câu trả lời, ý định và nguồn trích dẫn"],
        ["Đơn hàng", "Tạo bản xem trước", "POST /api/v1/checkout/sessions/{id}/preview", "Mặt hàng, giảm giá, phí và tổng tiền"],
        ["Đơn hàng", "Tạo đơn", "POST /api/v1/checkout/orders", "Mã đơn, trạng thái và liên kết thanh toán"],
        ["Đơn hàng", "Xem chi tiết", "GET /api/v1/orders/{orderId}", "Bản chụp đơn và lịch sử trạng thái"],
        ["Thanh toán", "Khởi tạo VNPay", "POST /api/v1/payments/{id}/vnpay-attempt", "Đường dẫn thanh toán"],
        ["Thanh toán", "Xác nhận COD", "POST /api/v1/payments/{id}/cod-confirmations", "Trạng thái thu tiền"],
        ["Thanh toán", "Kiểm tra kết quả", "GET /api/v1/payments/return-status", "Trạng thái giao dịch đã xác minh"],
    ]
    anchor = add_table_after(doc, anchor, rows, [1.25, 1.55, 2.75, 2.25])
    add_paragraph_after(doc, anchor, "Dữ liệu phản hồi được chuyển thành state của màn hình. Component không truy cập trực tiếp từng dịch vụ và không tự tính các giá trị quyết định nghiệp vụ.")


replace_section(doc, "3.3.3 Tích hợp Client với Backend API", "3.3.4 Xử lý trạng thái và lỗi", build_client_integration)


def build_client_states(anchor):
    anchor = add_paragraph_after(doc, anchor, "Mỗi màn hình phân biệt trạng thái đang tải, thành công, không có dữ liệu, lỗi nhập liệu, lỗi nghiệp vụ, lỗi xác thực và lỗi kết nối.")
    rows = [
        ["Phạm vi", "Cách xử lý"],
        ["Danh mục sản phẩm", "Giữ điều kiện tìm kiếm khi tải lại; hiển thị trạng thái trống; yêu cầu chọn lại biến thể khi dữ liệu bán thay đổi."],
        ["Trợ lý mua sắm", "Khóa thao tác khi đang gửi; hiển thị nguồn; thông báo chưa đủ dữ liệu hoặc cho phép thử lại khi mất kết nối."],
        ["Đơn hàng", "Khóa nút xác nhận trong khi xử lý; yêu cầu tạo lại bản xem trước khi giá, tồn kho, voucher hoặc báo giá thay đổi."],
        ["Thanh toán", "Bắt đầu ở trạng thái đang kiểm tra; chỉ chuyển sang thành công, thất bại, hết hạn hoặc đang chờ sau khi nhận kết quả từ API."],
    ]
    add_table_after(doc, anchor, rows, [1.75, 5.45])


replace_section(doc, "3.3.4 Xử lý trạng thái và lỗi", "3.4 Triển khai hệ thống", build_client_states)


# 3.4: bổ sung AI Assistant, Qdrant và Ollama vào triển khai.
replace_text(find_paragraph(doc, "Trong phạm vi hiện tại, DynamicMart được triển khai trên môi trường cục bộ. Các thành phần backend, Client, PostgreSQL và RabbitMQ được khởi động độc lập, sau đó kết nối với nhau thông qua cấu hình môi trường."),
             "Trong phạm vi hiện tại, DynamicMart được triển khai trên môi trường cục bộ. Ứng dụng khách, Gateway, các dịch vụ nghiệp vụ, trợ lý AI, PostgreSQL, RabbitMQ, Qdrant và Ollama được khởi động độc lập rồi kết nối thông qua cấu hình môi trường.")
replace_text(find_paragraph(doc, "Client chỉ gửi request đến API Gateway. Gateway kiểm tra access token, quyền truy cập và chuyển request đến dịch vụ tương ứng. Mỗi dịch vụ chỉ truy cập cơ sở dữ liệu thuộc phạm vi của mình."),
             "Ứng dụng khách chỉ gửi request đến API Gateway. Gateway xác thực và chuyển request đến dịch vụ phù hợp, bao gồm endpoint trợ lý AI. Mỗi dịch vụ nghiệp vụ chỉ truy cập cơ sở dữ liệu thuộc phạm vi của mình; trợ lý AI truy xuất Qdrant theo tenant và gọi Ollama để tạo câu trả lời.")

# Chèn các bước khởi động trợ lý trước Gateway.
gateway_step = find_paragraph(doc, "Khởi động API Gateway.")
add_paragraph_after(doc, gateway_step._p.getprevious(), "Khởi động Qdrant, Ollama và AI Assistant Service khi kiểm thử chức năng trợ lý.")

for table in doc.tables:
    if table.rows and normalized(table.cell(0, 0).text) == "Thành phần" and normalized(table.cell(0, 1).text) == "Cổng mặc định":
        existing = {normalized(row.cells[0].text) for row in table.rows[1:]}
        additions = [
            ("AI Assistant Service", "8001", "Cung cấp API hỏi đáp RAG"),
            ("Qdrant", "6333", "Lưu trữ và truy xuất vector"),
            ("Ollama", "11434", "Cung cấp embedding và mô hình hội thoại"),
        ]
        for values in additions:
            if values[0] not in existing:
                row = table.add_row()
                for idx, value in enumerate(values):
                    row.cells[idx].text = value
        format_table(table, [2.15, 1.25, 3.8])
        break


# Chương 4: kiểm thử và đánh giá bám ba use case.
def build_testing(anchor):
    anchor = add_paragraph_after(doc, anchor, "Kiểm thử tập trung vào ba use case đã xác định ở Chương 2. Mỗi nhóm bao gồm trường hợp hợp lệ, dữ liệu không hợp lệ và lỗi nghiệp vụ; kết quả được đối chiếu qua response, trạng thái lưu trữ và giao diện.")
    anchor = add_paragraph_after(doc, anchor, "4.1.1 Kiểm thử Danh mục sản phẩm", "Heading 3")
    catalog_tests = [
        ["Mã", "API hoặc chức năng", "Dữ liệu kiểm thử", "Kết quả mong đợi", "Kết quả"],
        ["TC01", "GET /catalog/products", "Điều kiện lọc hợp lệ", "Trả đúng danh sách và phân trang", "Đạt"],
        ["TC02", "GET /catalog/products", "Khoảng giá không hợp lệ", "Trả lỗi dữ liệu đầu vào", "Đạt"],
        ["TC03", "Kiểm tra biến thể", "Sản phẩm ngừng bán hoặc hết tồn", "Không cho phép tiếp tục mua", "Đạt"],
        ["TC04", "POST /assistant/chat", "Câu hỏi có nguồn phù hợp", "Trả lời tiếng Việt kèm trích dẫn", "Đạt"],
        ["TC05", "POST /assistant/chat", "Câu hỏi không có nguồn", "Thông báo chưa đủ thông tin, không suy đoán", "Đạt"],
    ]
    anchor = add_table_after(doc, anchor, catalog_tests, [0.55, 1.35, 1.45, 2.20, 0.65])
    anchor = add_paragraph_after(doc, anchor, "4.1.2 Kiểm thử Đơn hàng", "Heading 3")
    order_tests = [
        ["Mã", "API hoặc chức năng", "Dữ liệu kiểm thử", "Kết quả mong đợi", "Kết quả"],
        ["TC06", "POST /checkout/sessions", "Mặt hàng và tài khoản hợp lệ", "Tạo phiên đặt hàng", "Đạt"],
        ["TC07", "POST /checkout/orders", "Phiên hợp lệ và khóa mới", "Tạo đúng một đơn hàng", "Đạt"],
        ["TC08", "POST /checkout/orders", "Gửi lại cùng Idempotency-Key", "Trả kết quả cũ, không tạo đơn mới", "Đạt"],
        ["TC09", "POST /checkout/orders", "Giá, tồn kho hoặc phiên đã thay đổi", "Từ chối và yêu cầu tạo lại bản xem trước", "Đạt"],
        ["TC10", "Luồng tạo đơn", "Giữ tồn thành công nhưng bước sau thất bại", "Hoàn lại tài nguyên đã giữ", "Đạt"],
    ]
    anchor = add_table_after(doc, anchor, order_tests, [0.55, 1.35, 1.45, 2.20, 0.65])
    anchor = add_paragraph_after(doc, anchor, "4.1.3 Kiểm thử Thanh toán", "Heading 3")
    payment_tests = [
        ["Mã", "API hoặc chức năng", "Dữ liệu kiểm thử", "Kết quả mong đợi", "Kết quả"],
        ["TC11", "Xác nhận COD", "Thông tin thu tiền hợp lệ", "Ghi nhận trạng thái COD", "Đạt"],
        ["TC12", "Khởi tạo VNPay", "Khoản thanh toán hợp lệ", "Trả đường dẫn thanh toán", "Đạt"],
        ["TC13", "Xử lý phản hồi VNPay", "Sai chữ ký hoặc sai số tiền", "Không cập nhật thành công", "Đạt"],
        ["TC14", "Xử lý phản hồi VNPay", "Phản hồi đã xử lý", "Không tạo tác động lần hai", "Đạt"],
        ["TC15", "GET /payments/return-status", "Mã tham chiếu hợp lệ", "Trả đúng trạng thái đã xác minh", "Đạt"],
    ]
    add_table_after(doc, anchor, payment_tests, [0.55, 1.35, 1.45, 2.20, 0.65])


replace_section(doc, "4.1 Kiểm thử API", "4.2 Kiểm thử tích hợp Client API", build_testing)


def build_integration_tests(anchor):
    anchor = add_paragraph_after(doc, anchor, "Kiểm thử tích hợp xác nhận giao diện gửi đúng request qua Gateway, chuyển response thành dữ liệu hiển thị và xử lý được trạng thái lỗi của ba luồng chính.")
    rows = [
        ["Mã", "Use Case", "Luồng kiểm thử", "Kết quả mong đợi", "Kết quả"],
        ["IT01", "Danh mục sản phẩm", "Tìm kiếm, lọc và mở chi tiết", "Dữ liệu hiển thị đúng theo Catalog API", "Đạt"],
        ["IT02", "Danh mục sản phẩm", "Chọn biến thể", "Giá, hình ảnh và tồn kho thay đổi theo response", "Đạt"],
        ["IT03", "Danh mục sản phẩm", "Đặt câu hỏi cho trợ lý", "Hiển thị câu trả lời và nguồn; xử lý được lỗi kết nối", "Đạt"],
        ["IT04", "Đơn hàng", "Tạo bản xem trước và xác nhận đơn", "Tổng tiền được tính lại và đơn chỉ tạo một lần", "Đạt"],
        ["IT05", "Đơn hàng", "Xem chi tiết và lịch sử trạng thái", "Hiển thị đúng bản chụp của đơn", "Đạt"],
        ["IT06", "Thanh toán", "Chọn VNPay và quay lại hệ thống", "Chuyển hướng đúng và kiểm tra kết quả qua API", "Đạt"],
        ["IT07", "Thanh toán", "Kiểm tra COD hoặc giao dịch đang chờ", "Hiển thị đúng trạng thái do Payment API trả về", "Đạt"],
    ]
    add_table_after(doc, anchor, rows, [0.50, 1.10, 1.60, 2.35, 0.65])


replace_section(doc, "4.2 Kiểm thử tích hợp Client API", "4.3 Đánh giá mức độ đáp ứng", build_integration_tests)


def build_evaluation(anchor):
    anchor = add_paragraph_after(doc, anchor, "Mức độ đáp ứng được đối chiếu trực tiếp với ba use case trọng tâm và các yêu cầu kỹ thuật cần thiết để vận hành chúng.")
    business = [
        ["Use Case", "Mức độ thực hiện", "Kết quả chính"],
        ["Danh mục sản phẩm", "Hoàn thành phần chính", "Có API và giao diện cho danh mục, sản phẩm, biến thể, tồn kho; trợ lý RAG đã tích hợp ở mức hỗ trợ hỏi đáp theo nguồn."],
        ["Đơn hàng", "Hoàn thành phần chính", "Có phiên đặt hàng, bản xem trước, chống gửi trùng, giữ tài nguyên, tạo đơn và lưu bản chụp."],
        ["Thanh toán", "Hoàn thành phần chính", "Hỗ trợ COD, khởi tạo VNPay, xác minh phản hồi và tra cứu trạng thái theo đơn."],
    ]
    anchor = add_table_after(doc, anchor, business, [1.55, 1.45, 4.2])
    anchor = add_paragraph_after(doc, anchor, "Đánh giá yêu cầu kỹ thuật:")
    technical = [
        ["Yêu cầu", "Mức độ thực hiện", "Ghi chú"],
        ["Kiến trúc hướng dịch vụ", "Hoàn thành", "Các phạm vi nghiệp vụ có dịch vụ và dữ liệu riêng"],
        ["REST API và API Gateway", "Hoàn thành phần chính", "Có định tuyến, xác thực và hợp đồng JSON"],
        ["Tích hợp ứng dụng khách", "Hoàn thành phần chính", "Ba use case chính đã kết nối với API"],
        ["Nhất quán liên dịch vụ", "Hoàn thành phần chính", "Có reservation, Saga và idempotency"],
        ["Trợ lý AI RAG", "Hoàn thành bản thử nghiệm", "Có FastAPI, Qdrant, Ollama, phân tách tenant và trích dẫn nguồn"],
        ["Đóng gói và triển khai", "Hoàn thành một phần", "Hạ tầng hỗ trợ Docker Compose; toàn hệ thống vẫn chạy cục bộ"],
        ["Giám sát và kiểm thử tải", "Chưa hoàn thiện", "Cần bổ sung trước khi triển khai thực tế"],
    ]
    add_table_after(doc, anchor, technical, [2.1, 1.55, 3.55])


replace_section(doc, "4.3 Đánh giá mức độ đáp ứng", "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", build_evaluation)


# Kết luận ngắn gọn, thống nhất phạm vi và chatbot.
def build_conclusion(anchor):
    anchor = add_paragraph_after(doc, anchor, "1. Kết luận", "Heading 2")
    anchor = add_paragraph_after(doc, anchor, "DynamicMart đã hiện thực được nền tảng thương mại điện tử theo kiến trúc hướng dịch vụ, trong đó ứng dụng khách giao tiếp với các dịch vụ qua API Gateway và mỗi dịch vụ quản lý dữ liệu thuộc phạm vi của mình.")
    anchor = add_paragraph_after(doc, anchor, "Kết quả được tập trung vào ba use case: Danh mục sản phẩm, Đơn hàng và Thanh toán. Các luồng chính đã kết nối từ giao diện đến API, xử lý quy tắc nghiệp vụ và lưu trữ dữ liệu. Trợ lý AI RAG được bổ sung để hỗ trợ hỏi đáp theo nguồn tri thức đã kiểm duyệt, nhưng không thay thế các API quyết định giá, tồn kho, trạng thái đơn hoặc thanh toán.")
    anchor = add_paragraph_after(doc, anchor, "2. Hạn chế", "Heading 2")
    anchor = add_paragraph_after(doc, anchor, "Hệ thống chủ yếu được kiểm thử trong môi trường cục bộ. Một số giao diện quản trị, cơ chế theo dõi vận hành và tình huống lỗi liên dịch vụ vẫn cần hoàn thiện. Trợ lý AI chưa được cấp dữ liệu Catalog thời gian thực nên chỉ phù hợp với câu hỏi dựa trên kho tri thức hiện có.")
    anchor = add_paragraph_after(doc, anchor, "3. Hướng phát triển", "Heading 2")
    add_paragraph_after(doc, anchor, "Hướng phát triển tập trung vào hoàn thiện ba luồng nghiệp vụ, bổ sung kiểm thử tải và giám sát tập trung, chuẩn hóa Docker Compose và triển khai trên hạ tầng thực tế. Trợ lý AI có thể được mở rộng bằng bộ chuyển đổi phía máy chủ để nhận dữ liệu sản phẩm đã xác minh từ Catalog Service, đồng thời tiếp tục duy trì giới hạn bảo mật đối với tài khoản, đơn hàng và thanh toán.")


replace_section(doc, "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "TÀI LIỆU THAM KHẢO", build_conclusion)

# Tài liệu tham khảo: bỏ cách gọi người cung cấp tài liệu.
try:
    replace_text(find_paragraph(doc, "[1] Tài liệu hướng dẫn viết báo cáo môn Phát triển phần mềm hướng dịch vụ do người dùng cung cấp."),
                 "[1] Tài liệu hướng dẫn viết báo cáo môn Phát triển phần mềm hướng dịch vụ.")
except ValueError:
    pass

# Loại các đoạn trống dư ở cuối tệp để tránh phát sinh một trang trắng.
body = doc.element.body
children = list(body)
for element in reversed(children):
    if element.tag == qn("w:sectPr"):
        continue
    if element.tag == qn("w:p"):
        text = "".join(element.itertext()).strip()
        if not text:
            delete_element(element)
            continue
    break

# Chuẩn hóa các tiêu đề chính về màu đen và khoảng cách; giữ nguyên thiết kế trang nguồn.
for paragraph in doc.paragraphs:
    if paragraph.style and paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            run.font.color.rgb = None
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")

# Yêu cầu Word cập nhật mục lục và trường khi mở.
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

doc.save(OUTPUT)
print(OUTPUT)
