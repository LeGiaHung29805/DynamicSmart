from pathlib import Path
import sys
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"E:\dynamicmart\docs\report-edit\Bao-cao-nguon.docx")
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(r"E:\dynamicmart\docs\report-edit\noi-dung-nguon.txt")


def iter_blocks(doc):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


doc = Document(SRC)
lines = []
for idx, block in enumerate(iter_blocks(doc), start=1):
    if isinstance(block, Paragraph):
        text = block.text.replace("\n", " ").strip()
        if text:
            lines.append(f"P{idx:04d}\t[{block.style.name}]\t{text}")
    else:
        lines.append(f"T{idx:04d}\t[TABLE {len(block.rows)}x{len(block.columns)}]")
        for r_idx, row in enumerate(block.rows, start=1):
            cells = [" / ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()) for cell in row.cells]
            lines.append(f"  R{r_idx:03d}\t" + " || ".join(cells))

OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
print(f"output={OUT}")
